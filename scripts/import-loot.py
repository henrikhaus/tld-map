"""Read Bashrobe's workbook without changing it; retain cells and comments as data."""
import json, re, sys
from pathlib import Path
import openpyxl
source = Path(sys.argv[1])
book = openpyxl.load_workbook(source, data_only=True)
records, comments, counts, groups = {}, [], {}, []
def normalized(text):
    return re.sub(r'\s+', ' ', text.replace('’', "'").strip()).lower()
def note_data(cell):
    note = cell.comment.text.strip() if cell.comment else ''
    community = note.startswith('[Threaded comment]')
    if community: note = note.split('Comment:', 1)[-1].strip()
    return dict(note=note, community=community)
for sheet in book:
    active_groups = {}
    set_no, region, count = int(sheet.title[-1]), None, 0
    for row in sheet.iter_rows(min_row=3):
        if row[0].value:
            region = str(row[0].value).strip()
            active_groups = {}
        for col in range(1, 8):
            cell = row[col]
            if not cell.value or not region: continue
            item, location = str(sheet.cell(1, col + 1).value), str(cell.value).strip()
            key = '|'.join(map(normalized, [region, item, location]))
            if key not in records:
                records[key] = dict(id=f'loot-{len(records)+1}', region=region, item=item, location=location, sets=[], sources=[])
            entry = records[key]
            if set_no not in entry['sets']: entry['sets'].append(set_no)
            previous = sheet.cell(cell.row - 1, cell.column)
            # Horizontal borders encode separate, independently available spawns.
            if col not in active_groups or getattr(cell.border.top, 'style', None) or getattr(previous.border.bottom, 'style', None):
                group = dict(id=f'{sheet.title}:{cell.coordinate}', set=set_no, region=region, item=item, entries=[], startCell=cell.coordinate, endCell=cell.coordinate)
                groups.append(group)
                active_groups[col] = group
            group = active_groups[col]
            group['entries'].append(entry['id'])
            group['endCell'] = cell.coordinate
            entry['sources'].append(dict(sheet=sheet.title, cell=cell.coordinate, group=group['id'], **note_data(cell)))
            count += 1
        for cell in row[1:8]:
            if cell.comment and not cell.value and region:
                comments.append(dict(sheet=sheet.title, set=set_no, cell=cell.coordinate, region=region, item=str(sheet.cell(1, cell.column).value), **note_data(cell)))
    counts[sheet.title] = count
result = dict(source=source.name, author='Bashrobe', version=None, counts=counts, entries=list(records.values()), additionalComments=comments, spawnGroups=groups)
Path('data/loot.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(dict(counts=counts, uniqueLocations=len(records), regions=sorted({r['region'] for r in records.values()}))))
